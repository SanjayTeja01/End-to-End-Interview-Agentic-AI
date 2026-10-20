"""Text extraction: PyMuPDF for PDF, python-docx for DOCX, UTF-8 for text/markdown."""
from __future__ import annotations

import io
from dataclasses import dataclass


@dataclass
class Extracted:
    text: str
    pages: int | None
    meta: dict


class ExtractionError(ValueError):
    pass


def extract_pdf(data: bytes) -> Extracted:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not open PDF") from exc
    if doc.needs_pass:
        raise ExtractionError("Encrypted PDFs are not supported")
    texts = [page.get_text("text") for page in doc]
    pages = doc.page_count
    doc.close()
    text = "\n\n".join(texts)
    if not text.strip():
        raise ExtractionError("PDF contains no extractable text (scanned documents are not supported)")
    return Extracted(text=text, pages=pages, meta={"pages": pages})


def extract_docx(data: bytes) -> Extracted:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("Could not open DOCX") from exc
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("DOCX contains no text")
    return Extracted(text=text, pages=None, meta={"paragraphs": len(d.paragraphs), "tables": len(d.tables)})


def extract(data: bytes, ext: str) -> Extracted:
    if ext == "pdf":
        return extract_pdf(data)
    if ext == "docx":
        return extract_docx(data)
    if ext in ("txt", "md"):
        return Extracted(text=data.decode("utf-8", errors="replace"), pages=None, meta={})
    raise ExtractionError(f"Unsupported extension {ext}")
